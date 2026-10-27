import io
import os
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_HEADERS = [
    "ID",
    "Full Name",
    "Job Title",
    "Company",
    "Phone Number",
    "Email Address",
    "Website",
    "Address",
    "Other Details",
    "Scanned At",
    "Raw Extracted Text"
]

# Column width constraints: (min_width, max_width)
COL_WIDTH_CONSTRAINTS = {
    1: (10, 14),   # ID
    2: (18, 30),   # Full Name
    3: (16, 30),   # Job Title
    4: (20, 35),   # Company
    5: (16, 24),   # Phone Number
    6: (22, 36),   # Email Address
    7: (20, 36),   # Website
    8: (25, 55),   # Address
    9: (20, 50),   # Other Details
    10: (20, 24),  # Scanned At
    11: (25, 65),  # Raw Extracted Text
}

def create_excel_workbook(cards: List[Dict[str, Any]]) -> Workbook:
    """
    Builds a beautifully styled, professional Excel workbook containing business cards.
    Includes:
      - Clean slate-navy header styling
      - High-contrast typography & auto-filtered columns
      - Preserved text formats for phone numbers (no scientific notation / lost leading zeros)
      - Alternating zebra row striping
      - Word wrapping on multi-line addresses and raw OCR text
      - Sticky frozen header row (Freeze Panes)
      - Perfectly calibrated column widths
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Business Cards"

    # Explicitly ensure gridlines are visible
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align_left = Alignment(horizontal="left", vertical="center", wrap_text=False)
    header_align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)

    zebra_even_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    zebra_odd_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    body_font = Font(name="Calibri", size=10.5, color="0F172A")
    body_font_mono = Font(name="Consolas", size=10, color="334155")
    link_font = Font(name="Calibri", size=10.5, color="2563EB", underline="single")

    thin_border_side = Side(border_style="thin", color="E2E8F0")
    header_border_side = Side(border_style="thin", color="334155")

    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )
    header_border = Border(
        left=header_border_side,
        right=header_border_side,
        top=header_border_side,
        bottom=header_border_side
    )

    # 1. Write Header Row
    ws.row_dimensions[1].height = 28
    for col_idx, header_title in enumerate(EXCEL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = header_border
        cell.alignment = header_align_center if col_idx == 1 else header_align_left

    # 2. Write Data Rows
    for row_idx, card in enumerate(cards, start=2):
        is_odd_data_row = (row_idx % 2 == 1)
        row_fill = zebra_odd_fill if is_odd_data_row else zebra_even_fill
        ws.row_dimensions[row_idx].height = 22

        # Extract values
        card_id = str(card.get("id") or "")
        name = str(card.get("name") or "")
        title = str(card.get("title") or "")
        company = str(card.get("company") or "")
        phone = str(card.get("phone") or "")
        email = str(card.get("email") or "")
        website = str(card.get("website") or "")
        address = str(card.get("address") or "")
        other_details = str(card.get("other_details") or "")
        scanned_at = str(card.get("scanned_at") or "")
        raw_text = str(card.get("raw_text") or "")

        row_values = [
            card_id,
            name,
            title,
            company,
            phone,
            email,
            website,
            address,
            other_details,
            scanned_at,
            raw_text
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = val
            cell.fill = row_fill
            cell.border = cell_border
            cell.font = body_font

            # Specific column alignments and types
            if col_idx == 1:
                # ID
                cell.font = body_font_mono
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = "@"
            elif col_idx == 5:
                # Phone: STRICT text format to prevent scientific notation or dropping leading 0/+
                cell.number_format = "@"
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 6 and email and "@" in email:
                # Email
                cell.alignment = Alignment(horizontal="left", vertical="center")
                if not email.startswith("mailto:"):
                    cell.hyperlink = f"mailto:{email}"
                    cell.font = link_font
            elif col_idx == 7 and website and "." in website:
                # Website
                cell.alignment = Alignment(horizontal="left", vertical="center")
                url = website if (website.startswith("http://") or website.startswith("https://")) else f"https://{website}"
                cell.hyperlink = url
                cell.font = link_font
            elif col_idx == 10:
                # Scanned At
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (8, 9, 11):
                # Address, Other Details, Raw Text - enable multiline wrap & top alignment
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                # If content has multiple newlines, adjust row height so it displays well
                newlines = val.count("\n")
                if newlines > 1:
                    dynamic_height = min(120, max(22, 18 * (newlines + 1)))
                    if ws.row_dimensions[row_idx].height < dynamic_height:
                        ws.row_dimensions[row_idx].height = dynamic_height
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # 3. Dynamic Column Widths
    total_cols = len(EXCEL_HEADERS)
    total_rows = len(cards) + 1

    for col_idx in range(1, total_cols + 1):
        col_letter = get_column_letter(col_idx)
        min_w, max_w = COL_WIDTH_CONSTRAINTS.get(col_idx, (15, 40))

        # Calculate max line length among header & data rows
        max_len = len(str(EXCEL_HEADERS[col_idx - 1]))
        for row_idx in range(2, total_rows + 1):
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val:
                # If multiline, take longest line rather than whole string
                lines = str(cell_val).split("\n")
                line_max = max(len(l) for l in lines) if lines else 0
                if line_max > max_len:
                    max_len = line_max

        # Apply padding and clamp between min and max
        computed_width = max(min_w, min(max_len + 4, max_w))
        ws.column_dimensions[col_letter].width = computed_width

    # 4. Freeze header row so it stays visible while scrolling
    ws.freeze_panes = "A2"

    # 5. Enable AutoFilter across all headers
    last_col_letter = get_column_letter(total_cols)
    ws.auto_filter.ref = f"A1:{last_col_letter}{total_rows}"

    return wb

def generate_excel_bytes(cards: List[Dict[str, Any]]) -> bytes:
    """Generates in-memory .xlsx file bytes for HTTP download."""
    wb = create_excel_workbook(cards)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

def write_cards_to_excel_file(cards: List[Dict[str, Any]], filepath: str) -> None:
    """Writes the full list of cards to an Excel file on disk."""
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb = create_excel_workbook(cards)
    wb.save(filepath)
